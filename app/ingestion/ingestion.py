from pathlib import Path
 
from docx import Document

from langchain_core.documents import Document as LangChainDocument

from langchain_text_splitters import RecursiveCharacterTextSplitter
 
from app.core.db import get_vector_store
 
COLLECTION_NAME = "smart_banking_kb"
 
 
def load_smart_banking_document(file_path: str) -> list[LangChainDocument]:

    path = Path(file_path)
 
    if not path.exists():

        raise FileNotFoundError(f"Knowledge base not found: {file_path}")
 
    source = Document(file_path)

    documents = []
 
    for index, paragraph in enumerate(source.paragraphs):

        text = paragraph.text.strip()

        if text:

            documents.append(

                LangChainDocument(

                    page_content=text,

                    metadata={

                        "document_name": path.name,

                        "source_page": None,

                        "chunk_type": "text",

                        "product_category": "banking",

                        "element_index": index,

                    },

                )

            )
 
    for table_index, table in enumerate(source.tables):

        rows = []

        for row in table.rows:

            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]

            rows.append(" | ".join(cells))
 
        if rows:

            documents.append(

                LangChainDocument(

                    page_content="\n".join(rows),

                    metadata={

                        "document_name": path.name,

                        "source_page": None,

                        "chunk_type": "table",

                        "product_category": "banking",

                        "element_index": table_index,

                    },

                )

            )
 
    return documents
 
 
def ingest_smart_banking_kb(file_path: str = "data/KB_Smart_Banking.docx") -> int:

    documents = load_smart_banking_document(file_path)
 
    splitter = RecursiveCharacterTextSplitter(

        chunk_size=1000,

        chunk_overlap=200,

        length_function=len,

    )
 
    chunks = splitter.split_documents(documents)
 
    for index, chunk in enumerate(chunks):

        chunk.metadata["chunk_id"] = index

        chunk.metadata["collection_name"] = COLLECTION_NAME
 
    vector_store = get_vector_store(COLLECTION_NAME)

    vector_store.add_documents(chunks)
 
    return len(chunks)
 
 
if __name__ == "__main__":

    count = ingest_smart_banking_kb()

    print(f"Ingestion completed. Chunks: {count}")